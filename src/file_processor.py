"""
File processor: extracts text from PDF, DOCX, PPTX, XLSX, CSV, TXT files.
"""

import io
import csv
from pathlib import Path
from typing import Tuple


def extract_text_from_file(uploaded_file) -> Tuple[str, str]:
    """
    Extract text from a Streamlit UploadedFile object.
    Returns (extracted_text, status_message).
    """
    filename = uploaded_file.name
    ext = Path(filename).suffix.lower().lstrip(".")
    raw_bytes = uploaded_file.read()

    try:
        if ext == "pdf":
            return _extract_pdf(raw_bytes, filename)
        elif ext in ("docx", "doc"):
            return _extract_docx(raw_bytes, filename)
        elif ext in ("pptx", "ppt"):
            return _extract_pptx(raw_bytes, filename)
        elif ext in ("xlsx", "xls"):
            return _extract_xlsx(raw_bytes, filename)
        elif ext == "csv":
            return _extract_csv(raw_bytes, filename)
        elif ext in ("txt", "md"):
            return _extract_text(raw_bytes, filename)
        else:
            return "", f"⚠️ {filename}: Unsupported file type (.{ext})"
    except Exception as e:
        return "", f"❌ {filename}: Could not process file — {str(e)}"


def _extract_pdf(raw_bytes: bytes, filename: str) -> Tuple[str, str]:
    try:
        import pdfplumber
        text_parts = []
        with pdfplumber.open(io.BytesIO(raw_bytes)) as pdf:
            for i, page in enumerate(pdf.pages):
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(f"[Page {i + 1}]\n{page_text}")
        full_text = "\n\n".join(text_parts)
        if not full_text.strip():
            return "", f"⚠️ {filename}: No extractable text found (may be a scanned PDF)."
        return full_text, f"✅ {filename} — {len(pdf.pages)} pages extracted"
    except Exception as e:
        return "", f"❌ {filename}: PDF extraction failed — {str(e)}"


def _extract_docx(raw_bytes: bytes, filename: str) -> Tuple[str, str]:
    try:
        from docx import Document
        doc = Document(io.BytesIO(raw_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        full_text = "\n\n".join(paragraphs)
        return full_text, f"✅ {filename} — {len(paragraphs)} paragraphs extracted"
    except Exception as e:
        return "", f"❌ {filename}: DOCX extraction failed — {str(e)}"


def _extract_pptx(raw_bytes: bytes, filename: str) -> Tuple[str, str]:
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(raw_bytes))
        slide_texts = []
        for i, slide in enumerate(prs.slides):
            parts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    parts.append(shape.text.strip())
            if parts:
                slide_texts.append(f"[Slide {i + 1}]\n" + "\n".join(parts))
        full_text = "\n\n".join(slide_texts)
        return full_text, f"✅ {filename} — {len(prs.slides)} slides extracted"
    except Exception as e:
        return "", f"❌ {filename}: PPTX extraction failed — {str(e)}"


def _extract_xlsx(raw_bytes: bytes, filename: str) -> Tuple[str, str]:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
        sheet_texts = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                row_str = " | ".join(str(c) for c in row if c is not None)
                if row_str.strip():
                    rows.append(row_str)
            if rows:
                sheet_texts.append(f"[Sheet: {sheet_name}]\n" + "\n".join(rows))
        full_text = "\n\n".join(sheet_texts)
        return full_text, f"✅ {filename} — {len(wb.sheetnames)} sheets extracted"
    except Exception as e:
        return "", f"❌ {filename}: XLSX extraction failed — {str(e)}"


def _extract_csv(raw_bytes: bytes, filename: str) -> Tuple[str, str]:
    try:
        decoded = raw_bytes.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(decoded))
        rows = [" | ".join(row) for row in reader if any(c.strip() for c in row)]
        full_text = "\n".join(rows)
        return full_text, f"✅ {filename} — {len(rows)} rows extracted"
    except Exception as e:
        return "", f"❌ {filename}: CSV extraction failed — {str(e)}"


def _extract_text(raw_bytes: bytes, filename: str) -> Tuple[str, str]:
    try:
        full_text = raw_bytes.decode("utf-8", errors="replace")
        return full_text, f"✅ {filename} — plain text extracted"
    except Exception as e:
        return "", f"❌ {filename}: Text extraction failed — {str(e)}"


def truncate_for_context(text: str, max_chars: int = 60000) -> str:
    """Truncate extracted text to fit within LLM context window."""
    if len(text) <= max_chars:
        return text
    half = max_chars // 2
    return (
        text[:half]
        + f"\n\n[... content truncated for context window — {len(text) - max_chars} characters omitted ...]\n\n"
        + text[-half:]
    )
